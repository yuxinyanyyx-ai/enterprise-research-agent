from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from .constant import OUTPUT_DIR


def export_multi_query_result(result: dict):
    """
    将多药物 DMF 查询结果导出为 Excel。

    result:
        search_multiple_ingredients() 返回的结果
    """

    if not result.get("success") and not result.get("results"):
        raise ValueError("没有可导出的查询结果。")

    # 确保 outputs 目录存在
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # 文件名带时间，避免覆盖之前结果
    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_path = (
        OUTPUT_DIR
        / f"dmf_results_{timestamp}.xlsx"
    )

    # ==========================
    # 创建 Excel
    # ==========================

    workbook = Workbook()

    # 默认 Sheet
    sheet = workbook.active
    sheet.title = "DMF查询结果"

    # ==========================
    # 表头
    # ==========================

    headers = [
        "序号",
        "查询关键词",
        "DMF编号",
        "申请商名称",
        "成分",
        "有效日期",
        "查询状态",
    ]

    sheet.append(headers)

    # 表头样式
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ==========================
    # 写入查询结果
    # ==========================

    row_number = 1

    for query_result in result["results"]:

        query = query_result.get(
            "query",
            {},
        )

        query_keyword = (
            query.get("ingredient")
            or query.get("dmf_no")
            or query.get("applicant_name")
            or ""
        )

        # ----------------------
        # 查询失败
        # ----------------------

        if not query_result.get("success"):
            sheet.append([
                row_number,
                query_keyword,
                "",
                "",
                "",
                "",
                query_result.get(
                    "message",
                    "查询失败",
                ),
            ])

            row_number += 1
            continue

        records = query_result.get(
            "records",
            [],
        )

        # ----------------------
        # 查询成功但 0 条
        # ----------------------

        if not records:
            sheet.append([
                row_number,
                query_keyword,
                "",
                "",
                "",
                "",
                "查询成功，无匹配记录",
            ])

            row_number += 1
            continue

        # ----------------------
        # 正常数据
        # ----------------------

        for record in records:
            sheet.append([
                row_number,
                query_keyword,
                record.get(
                    "dmf_no",
                    "",
                ),
                record.get(
                    "applicant_name",
                    "",
                ),
                record.get(
                    "ingredient",
                    "",
                ),
                record.get(
                    "valid_date",
                    "",
                ),
                "查询成功",
            ])

            row_number += 1

    # ==========================
    # Excel 基础格式
    # ==========================

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        sheet.dimensions
    )

    widths = {
        "A": 8,
        "B": 20,
        "C": 20,
        "D": 28,
        "E": 35,
        "F": 15,
        "G": 22,
    }

    for column, width in widths.items():
        sheet.column_dimensions[
            column
        ].width = width

    # 内容自动换行
    for row in sheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    # ==========================
    # 增加汇总 Sheet
    # ==========================

    summary = workbook.create_sheet(
        "查询汇总"
    )

    summary_data = [
        ["项目", "结果"],
        [
            "总体状态",
            result.get("message", ""),
        ],
        [
            "查询药物数",
            result.get("query_count", 0),
        ],
        [
            "成功查询数",
            result.get("success_count", 0),
        ],
        [
            "失败查询数",
            result.get("failed_count", 0),
        ],
        [
            "DMF记录总数",
            result.get("total_records", 0),
        ],
        [
            "导出时间",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ],
    ]

    for row in summary_data:
        summary.append(row)

    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font

    summary.column_dimensions[
        "A"
    ].width = 20

    summary.column_dimensions[
        "B"
    ].width = 30

    # ==========================
    # 保存
    # ==========================

    workbook.save(output_path)

    return output_path