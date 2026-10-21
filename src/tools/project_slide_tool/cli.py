from __future__ import annotations

import argparse
import json
from pathlib import Path

from .fill import run_fill
from .lpt import extract_bottom_tables
from .milestone import format_milestone_markdown, parse_milestone_slide
from .models import FillConfig
from .pptx_utils import read_table

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "Example.pptx"
DEFAULT_TEMPLATE = ROOT / "project_slide_template.pptx"


def cmd_create_template(args: argparse.Namespace) -> None:
    import sys

    sys.path.insert(0, str(ROOT))
    from create_template import build_template

    out = Path(args.output) if args.output else DEFAULT_TEMPLATE
    path = build_template(out)
    print(f"已生成模板: {path}")
    print("  格式与 Example.pptx 第 1 页相同（Office Timeline）")


def cmd_fill(args: argparse.Namespace) -> None:
    template = Path(args.template or DEFAULT_TEMPLATE)
    if not template.is_file() and EXAMPLE.is_file():
        print(f"模板不存在，正在从 Example.pptx 生成: {DEFAULT_TEMPLATE}")
        import sys

        sys.path.insert(0, str(ROOT))
        from create_template import build_template

        build_template()
        template = DEFAULT_TEMPLATE

    data = Path(args.data)
    output = Path(args.output or ROOT / "output" / "project_slide.pptx")
    cfg = FillConfig(
        template=str(template),
        data=str(data),
        output=str(output),
        data_lpt_slide=args.lpt_slide,
        data_milestone_slide=args.milestone_slide,
    )
    result = run_fill(cfg)
    print(f"Output: {result.output_path}")
    print(f"  Upper: {result.upper_milestones}")
    print(f"  Lower: {result.lower_milestones}")
    print(f"  Bottom source: {result.bottom_source}")
    print(f"  Risks filled: {result.bottom_risks_count}")
    skipped = result.skipped_milestones or {}
    for track in ("upper", "lower"):
        items = skipped.get(track) or []
        if items:
            print(f"  Error: {track} track could not place: {items}")


def cmd_extract_milestone(args: argparse.Namespace) -> None:
    from pptx import Presentation

    prs = Presentation(args.data)
    slide = prs.slides[args.slide - 1]
    for shape in slide.shapes:
        if shape.has_table and len(shape.table.columns) >= 9:
            table = read_table(shape)
            print(format_milestone_markdown(table))
            upper, lower = parse_milestone_slide(slide)
            print("\n## Parsed Upper (Global)")
            print(json.dumps(upper.__dict__, ensure_ascii=False, indent=2, default=lambda o: o.__dict__ if hasattr(o, "__dict__") else str(o)))
            print("\n## Parsed Lower (China)")
            print(json.dumps(lower.__dict__, ensure_ascii=False, indent=2, default=lambda o: o.__dict__ if hasattr(o, "__dict__") else str(o)))
            return
    raise SystemExit(f"第 {args.slide} 页未找到 milestone 表")


def cmd_extract_lpt(args: argparse.Namespace) -> None:
    data = Path(args.data)
    bt, source = extract_bottom_tables(data, args.slide - 1)
    print(f"来源: {source}")
    print(f"Global Status:\n{bt.global_status}\n")
    print(f"Global Next:\n{bt.global_next}\n")
    print(f"China Status:\n{bt.china_status}\n")
    print(f"China Next:\n{bt.china_next}\n")
    print(f"Risks ({len(bt.risks)}):")
    for risk, mit in bt.risks:
        print(f"  - {risk[:80]}")
        if mit:
            print(f"    → {mit[:80]}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Project Slide 填充工具（Example.pptx 格式）")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_tpl = sub.add_parser("create-template", help="从 Example.pptx 生成空白模板（同格式）")
    p_tpl.add_argument("-o", "--output", help="输出路径（默认 project_slide_template.pptx）")
    p_tpl.set_defaults(func=cmd_create_template)

    p_fill = sub.add_parser("fill", help="用数据文件第2/3页填充模板第1页")
    p_fill.add_argument("-d", "--data", required=True, help="数据 pptx（含第2、3页）")
    p_fill.add_argument("-t", "--template", help="模板（默认 project_slide_template.pptx，与 Example 同格式）")
    p_fill.add_argument("-o", "--output", help="输出路径")
    p_fill.add_argument("--lpt-slide", type=int, default=2)
    p_fill.add_argument("--milestone-slide", type=int, default=3)
    p_fill.set_defaults(func=cmd_fill)

    p_ms = sub.add_parser("extract-milestone", help="查看 milestone 解析结果")
    p_ms.add_argument("-d", "--data", required=True)
    p_ms.add_argument("--slide", type=int, default=3)
    p_ms.set_defaults(func=cmd_extract_milestone)

    p_lpt = sub.add_parser("extract-lpt", help="查看第2页底部表格 LLM 提取结果")
    p_lpt.add_argument("-d", "--data", required=True)
    p_lpt.add_argument("--slide", type=int, default=2)
    p_lpt.set_defaults(func=cmd_extract_lpt)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
