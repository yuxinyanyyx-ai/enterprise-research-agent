from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from .lpt import apply_bottom_tables, extract_bottom_tables
from .milestone import extract_slide_title, parse_milestone_slide
from .models import FillConfig, FillResult
from .pptx_zip_patch import (
    PatchCollector,
    apply_patches_to_slide_xml,
    copy_pptx,
    read_slide_xml,
    write_slide_xml,
)
from .otl_shapes import scan_track_clusters
from .timeline import apply_slide_title, apply_timeline, ensure_cluster_capacity


def run_fill(cfg: FillConfig) -> FillResult:
    template = Path(cfg.template)
    data = Path(cfg.data)
    output = Path(cfg.output)

    copy_pptx(template, output)
    slide_xml = read_slide_xml(output, 0)

    data_prs = Presentation(str(data))
    if len(data_prs.slides) < max(cfg.data_lpt_slide, cfg.data_milestone_slide):
        raise ValueError(f"数据文件至少需要 {max(cfg.data_lpt_slide, cfg.data_milestone_slide)} 页")

    if cfg.upper is not None and cfg.lower is not None:
        upper, lower = cfg.upper, cfg.lower
    else:
        upper, lower = parse_milestone_slide(data_prs.slides[cfg.data_milestone_slide - 1])

    if cfg.slide_title is not None:
        slide_title = cfg.slide_title
    else:
        slide_title = extract_slide_title(data_prs.slides[cfg.data_milestone_slide - 1])

    tpl_prs = Presentation(str(template))
    _, clusters = scan_track_clusters(tpl_prs.slides[0])
    slide_xml = ensure_cluster_capacity(slide_xml, clusters, upper, lower)
    write_slide_xml(output, 0, slide_xml)

    prs = Presentation(str(output))
    slide = prs.slides[0]
    collector = PatchCollector()

    if slide_title:
        apply_slide_title(slide, slide_title, collector)
    timeline = apply_timeline(slide, upper, lower, collector)

    if cfg.bottom is not None:
        bottom = cfg.bottom
        bottom_source = "provided"
    else:
        bottom, bottom_source = extract_bottom_tables(data, cfg.data_lpt_slide - 1)
    apply_bottom_tables(slide, bottom, collector)

    patched = apply_patches_to_slide_xml(slide_xml, collector)
    write_slide_xml(output, 0, patched)

    return FillResult(
        output_path=str(output),
        upper_milestones=timeline.get("applied", {}).get("upper", []),
        lower_milestones=timeline.get("applied", {}).get("lower", []),
        bottom_source=bottom_source,
        bottom_risks_count=len(bottom.risks),
        skipped_milestones=timeline.get("skipped", {}),
    )
