from __future__ import annotations

from datetime import datetime

from pptx.util import Pt

from .constants import (
    DATE_BOX_WIDTH,
    DATE_FONT_SIZE,
    LOWER_STAR_Y,
    MILESTONE_SYMBOL_KIND,
    MIN_SEGMENT_WIDTH,
    MIN_TRACK_GAP,
    NOTE_FONT_SIZE,
    SEGMENT_END_MARGIN,
    SYMBOL_DIAMOND,
    SYMBOL_FLAG,
    SYMBOL_STAR,
    SYMBOL_TRIANGLE,
    UPPER_STAR_Y,
)
from .models import MilestoneEvent, TrackData
from .otl_shapes import OTL_NOTE, otl_kind, scan_track_clusters
from .pptx_zip_patch import PatchCollector
from .pptx_utils import detect_track_threshold, find_slide_title_shape, set_shape_text, track_of
from .timeline_layout import (
    _build_quarter_labels,
    _quarters_for_range,
    col_center_x,
    compute_year_range,
    date_to_col,
    display_label,
    parse_date,
    read_grid_geometry,
    update_quarter_grid,
)

NOTE_PT = Pt(NOTE_FONT_SIZE)
DATE_PT = Pt(DATE_FONT_SIZE)


def _write_text(shape, text: str, collector: PatchCollector | None, **kwargs) -> None:
    if collector:
        font_size = kwargs.get("font_size")
        font_size_pt = int(font_size.pt) if font_size is not None and hasattr(font_size, "pt") else None
        collector.set_text(shape, text, font_size_pt=font_size_pt)
    else:
        set_shape_text(shape, text, **kwargs)


def _hide_cluster(cl: dict, collector: PatchCollector | None) -> None:
    for key in ("label", "date", "note", "title"):
        sh = cl.get(key)
        if sh:
            _write_text(sh, "", collector)


def _has_note(ev: MilestoneEvent) -> bool:
    return bool(ev.note and ev.note.strip())


def _milestones_for_track(data: TrackData) -> list[str]:
    items = [ms for ms in data.event_order if ms in data.events]

    def sort_key(ms: str) -> tuple:
        ev = data.events[ms]
        dt = parse_date(ev.date)
        idx = data.event_order.index(ms)
        if dt:
            return (0, dt, idx)
        return (1, datetime.max, idx)

    return sorted(items, key=sort_key)


def _grid_marker_x(dt: datetime, grid: dict, quarters: list) -> int:
    return col_center_x(grid, date_to_col(dt, quarters))


def _hide_note_block(note, collector: PatchCollector | None) -> None:
    if not note:
        return
    _write_text(note, "", collector)
    if collector:
        collector.hide(note)


def _symbol_kind_for_milestone(ms: str) -> str:
    return MILESTONE_SYMBOL_KIND.get(ms, SYMBOL_DIAMOND)


def _is_small_triangle(shape) -> bool:
    name = shape.name or ""
    return (
        "OTLSHAPE_M_" in name
        and "_Shape" in name
        and "_Title" not in name
        and shape.width < 200_000
    )


def _is_diamond_marker(shape) -> bool:
    name = shape.name or ""
    return name.startswith("Star:") or ("4 Points" in name and "Star" in name)


def _is_nda_star(shape) -> bool:
    return "五角星" in (shape.name or "")


def _is_flag_graphic(shape) -> bool:
    name = shape.name or ""
    return name.startswith("Graphic") and getattr(shape, "width", 999999) < 500_000


def _is_timeline_end_cap(shape) -> bool:
    return "Isosceles Triangle" in (shape.name or "")


def _is_track_marker(shape) -> bool:
    return (
        _is_small_triangle(shape)
        or _is_diamond_marker(shape)
        or _is_nda_star(shape)
        or _is_flag_graphic(shape)
    )


def _marker_shape_kind(shape) -> str | None:
    if _is_small_triangle(shape):
        return SYMBOL_TRIANGLE
    if _is_diamond_marker(shape):
        return SYMBOL_DIAMOND
    if _is_nda_star(shape):
        return SYMBOL_STAR
    if _is_flag_graphic(shape):
        return SYMBOL_FLAG
    return None


def _place_shape_center(shape, center_x: int, collector: PatchCollector | None) -> None:
    try:
        left = int(center_x - shape.width // 2)
        if collector:
            collector.set_left(shape, left)
        else:
            shape.left = left
    except Exception:
        pass


def _hide_shape(shape, collector: PatchCollector | None) -> None:
    if collector:
        collector.hide(shape)


def _collect_marker_pools(slide, track: str, threshold: int) -> dict[str, list]:
    star_y = UPPER_STAR_Y if track == "upper" else LOWER_STAR_Y
    pools: dict[str, list] = {
        SYMBOL_TRIANGLE: [],
        SYMBOL_DIAMOND: [],
        SYMBOL_STAR: [],
        SYMBOL_FLAG: [],
    }
    for shape in slide.shapes:
        kind = _marker_shape_kind(shape)
        if not kind:
            continue
        if track_of(shape, threshold) != track:
            continue
        if kind in (SYMBOL_DIAMOND, SYMBOL_STAR) and abs(shape.top - star_y) > 300_000:
            continue
        pools[kind].append(shape)
    for shapes in pools.values():
        shapes.sort(key=lambda s: s.left)
    return pools


def _reposition_symbols(
    slide,
    upper_placed: list[tuple[str, dict, int]],
    lower_placed: list[tuple[str, dict, int]],
    grid: dict,
    collector: PatchCollector | None,
) -> None:
    """Place milestone markers by name: triangle / diamond / star / flag."""
    threshold = detect_track_threshold(slide)
    grid_left = int(grid["left"])
    grid_right = int(grid["left"] + grid["width"])

    def place_track_symbols(track: str, placed: list[tuple[str, dict, int]]) -> None:
        pools = _collect_marker_pools(slide, track, threshold)
        used: set[int] = set()
        for ms, _cl, marker_x in placed:
            kind = _symbol_kind_for_milestone(ms)
            available = [s for s in pools[kind] if id(s) not in used]
            if not available:
                continue
            shape = available[0]
            _place_shape_center(shape, marker_x, collector)
            used.add(id(shape))

        for shapes in pools.values():
            for shape in shapes:
                if id(shape) not in used:
                    _hide_shape(shape, collector)

    place_track_symbols("upper", upper_placed)
    place_track_symbols("lower", lower_placed)

    for shape in slide.shapes:
        if not _is_timeline_end_cap(shape):
            continue
        if track_of(shape, threshold) != "upper":
            continue
        name = shape.name or ""
        try:
            if "128" in name:
                _place_shape_center(shape, grid_right, collector)
            else:
                _place_shape_center(shape, grid_left, collector)
        except Exception:
            pass


def _collect_note_pool(slide, track: str, threshold: int) -> list:
    pool: list = []
    for shape in slide.shapes:
        if otl_kind(shape, threshold) != OTL_NOTE:
            continue
        if track_of(shape, threshold) != track:
            continue
        pool.append(shape)
    pool.sort(key=lambda s: s.left)
    return pool


def _resolve_track_notes(
    placed: list[tuple[str, dict, int]],
    data: TrackData,
    slide,
    track: str,
    threshold: int,
    collector: PatchCollector | None,
) -> list[tuple[str, dict, int]]:
    """Assign colored note blocks by milestone need, not fixed cluster pairing."""
    pool = _collect_note_pool(slide, track, threshold)
    used: set[int] = set()
    resolved: list[tuple[str, dict, int]] = []

    for ms, cl, marker_x in placed:
        ev = data.events[ms]
        cl = dict(cl)
        cluster_note = cl.get("note")
        if _has_note(ev):
            pick = None
            if cluster_note is not None and id(cluster_note) not in used:
                pick = cluster_note
            else:
                available = [s for s in pool if id(s) not in used]
                if available:
                    pick = min(available, key=lambda s: abs(s.left - marker_x))
            if pick is not None:
                cl["note"] = pick
                used.add(id(pick))
            else:
                cl["note"] = None
        else:
            if cluster_note is not None:
                _hide_note_block(cluster_note, collector)
            cl["note"] = None
        resolved.append((ms, cl, marker_x))

    for shape in pool:
        if id(shape) not in used:
            _hide_note_block(shape, collector)
    return resolved


def _assign_data_to_clusters(
    clusters: list[dict],
    data: TrackData,
    grid: dict,
    quarters: list,
    collector: PatchCollector | None,
) -> dict[str, dict]:
    sorted_clusters = sorted(clusters, key=lambda c: c["anchor"])
    for cl in sorted_clusters:
        _hide_cluster(cl, collector)

    ranked: list[tuple[str, int]] = []
    for ms in _milestones_for_track(data):
        ev = data.events[ms]
        dt = parse_date(ev.date)
        if not dt:
            continue
        ranked.append((ms, _grid_marker_x(dt, grid, quarters)))
    ranked.sort(key=lambda t: t[1])

    available = list(sorted_clusters)
    mapping: dict[str, dict] = {}
    for ms, marker_x in ranked:
        if not available:
            break
        pick = min(available, key=lambda c: abs(c["anchor"] - marker_x))
        available.remove(pick)
        mapping[ms] = pick
    return mapping


def _build_placed(
    mapping: dict[str, dict],
    data: TrackData,
    grid: dict,
    quarters: list,
) -> list[tuple[str, dict, int]]:
    items: list[tuple[str, dict, int]] = []
    for ms in _milestones_for_track(data):
        cl = mapping.get(ms)
        ev = data.events.get(ms)
        if not cl or not ev or not ev.date:
            continue
        dt = parse_date(ev.date)
        if not dt:
            continue
        items.append((ms, cl, _grid_marker_x(dt, grid, quarters)))
    return sorted(items, key=lambda t: t[2])


def _position_label(cl: dict, marker_x: int, collector: PatchCollector | None) -> None:
    for key in ("label", "title"):
        sh = cl.get(key)
        if not sh or sh is cl.get("note"):
            continue
        try:
            left = int(marker_x - sh.width // 2)
            if collector:
                collector.set_left(sh, left)
            else:
                sh.left = left
        except Exception:
            pass


def _apply_segment_layout(
    placed: list[tuple[str, dict, int]],
    grid: dict,
    data: TrackData,
    collector: PatchCollector | None,
) -> None:
    """Date/label at marker; colored note block only when slide 3 has note text."""
    if not placed:
        return
    track_right = int(grid["left"] + grid["width"] - SEGMENT_END_MARGIN)

    for i, (ms, cl, marker_x) in enumerate(placed):
        ev = data.events[ms]
        has_note = _has_note(ev)

        date_sh = cl.get("date")
        if date_sh:
            try:
                width = max(int(date_sh.width), DATE_BOX_WIDTH)
                left = int(marker_x - width // 2)
                if collector:
                    collector.set_width(date_sh, width)
                    collector.set_left(date_sh, left)
                else:
                    date_sh.width = width
                    date_sh.left = left
            except Exception:
                pass

        _position_label(cl, marker_x, collector)

        note = cl.get("note")
        if has_note and note:
            if i + 1 < len(placed):
                available = int(placed[i + 1][2] - marker_x - MIN_TRACK_GAP)
            else:
                available = int(track_right - marker_x)
            seg_width = (
                max(MIN_SEGMENT_WIDTH, available)
                if available >= MIN_SEGMENT_WIDTH
                else max(80_000, available)
            )
            try:
                if collector:
                    collector.set_left(note, marker_x)
                    collector.set_width(note, seg_width)
                    collector.show(note)
                else:
                    note.left = marker_x
                    note.width = seg_width
                if note.has_text_frame:
                    note.text_frame.word_wrap = True
            except Exception:
                pass
        else:
            _hide_note_block(note, collector)


def _position_tracks(
    slide,
    upper_map: dict[str, dict],
    lower_map: dict[str, dict],
    upper: TrackData,
    lower: TrackData,
    grid: dict,
    quarters: list,
    collector: PatchCollector | None,
) -> tuple[list, list]:
    threshold = detect_track_threshold(slide)
    upper_placed = _build_placed(upper_map, upper, grid, quarters)
    lower_placed = _build_placed(lower_map, lower, grid, quarters)
    upper_placed = _resolve_track_notes(upper_placed, upper, slide, "upper", threshold, collector)
    lower_placed = _resolve_track_notes(lower_placed, lower, slide, "lower", threshold, collector)
    _apply_segment_layout(upper_placed, grid, upper, collector)
    _apply_segment_layout(lower_placed, grid, lower, collector)
    return upper_placed, lower_placed


def _set_cluster_label(cl: dict, ms: str, collector: PatchCollector | None) -> None:
    text = display_label(ms)
    if ms == "CTA Sub" and cl.get("title"):
        _write_text(cl["title"], text, collector)
        if cl.get("label"):
            _write_text(cl["label"], "", collector)
    elif cl.get("label"):
        _write_text(cl["label"], text, collector)
    elif cl.get("title"):
        _write_text(cl["title"], text, collector)


def apply_slide_title(slide, title: str, collector: PatchCollector | None = None) -> None:
    shape = find_slide_title_shape(slide)
    if shape and title:
        _write_text(shape, title, collector)


def count_dated_milestones(data: TrackData) -> int:
    n = 0
    for ms in _milestones_for_track(data):
        if parse_date(data.events[ms].date):
            n += 1
    return n


def _cluster_clone_source(clusters: list[dict]) -> dict | None:
    for cl in reversed(clusters):
        if cl.get("date") and (cl.get("label") or cl.get("title")):
            return cl
    return None


def ensure_cluster_capacity(
    slide_xml: str,
    clusters: dict[str, list[dict]],
    upper: TrackData,
    lower: TrackData,
) -> str:
    """Clone date/label slots when milestone count exceeds template clusters."""
    from .pptx_zip_patch import expand_track_clusters_in_xml

    for track, data in (("upper", upper), ("lower", lower)):
        need = count_dated_milestones(data)
        have = len(clusters[track])
        extra = need - have
        if extra <= 0:
            continue
        src = _cluster_clone_source(clusters[track])
        if not src:
            raise RuntimeError(f"{track} 轨需要 {need} 个里程碑槽位，模板只有 {have} 个且无法克隆")
        slide_xml = expand_track_clusters_in_xml(slide_xml, src, extra)
    return slide_xml


def apply_timeline(
    slide, upper: TrackData, lower: TrackData, collector: PatchCollector | None = None
) -> dict[str, list[str]]:
    applied: dict[str, list[str]] = {"upper": [], "lower": []}
    skipped: dict[str, list[str]] = {"upper": [], "lower": []}
    study, all_clusters = scan_track_clusters(slide)

    start_y, end_y = compute_year_range(upper, lower)
    update_quarter_grid(slide, start_y, end_y, collector)

    grid = read_grid_geometry(slide)
    quarters = (
        _quarters_for_range(start_y, end_y, grid["cols"])
        if collector is not None
        else _build_quarter_labels(slide)
    )
    upper_map = _assign_data_to_clusters(all_clusters["upper"], upper, grid, quarters, collector)
    lower_map = _assign_data_to_clusters(all_clusters["lower"], lower, grid, quarters, collector)
    upper_placed, lower_placed = _position_tracks(
        slide, upper_map, lower_map, upper, lower, grid, quarters, collector
    )
    upper_by_ms = {ms: cl for ms, cl, _ in upper_placed}
    lower_by_ms = {ms: cl for ms, cl, _ in lower_placed}

    def apply_track(track: str, data: TrackData, mapping: dict[str, dict]) -> None:
        if track in study:
            if data.study_label:
                _write_text(
                    study[track],
                    data.study_label.replace(" | ", "\n"),
                    collector,
                    force_small=True,
                )
            else:
                _write_text(study[track], "", collector)

        for ms in _milestones_for_track(data):
            ev = data.events[ms]
            cl = mapping.get(ms)
            if not cl:
                if ev.date:
                    skipped[track].append(ms)
                continue
            applied[track].append(ms)

            if cl.get("date") and ev.date:
                _write_text(cl["date"], ev.date, collector, font_size=DATE_PT)
            elif cl.get("date"):
                _write_text(cl["date"], "", collector)

            if _has_note(ev) and cl.get("note"):
                _write_text(cl["note"], ev.note, collector, font_size=NOTE_PT)
            else:
                _hide_note_block(cl.get("note"), collector)

            _set_cluster_label(cl, ms, collector)

    apply_track("upper", upper, upper_by_ms)
    apply_track("lower", lower, lower_by_ms)
    _reposition_symbols(slide, upper_placed, lower_placed, grid, collector)

    for shape in slide.shapes:
        if shape.has_text_frame and "Project slide" in shape.text_frame.text:
            _write_text(shape, "", collector)

    return {"applied": applied, "skipped": skipped}
