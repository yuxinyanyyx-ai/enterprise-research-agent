"""从 sources 下的 Office/PDF 文件抽取 chunk。"""

from __future__ import annotations

import hashlib
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from src.pec.knowledge.markdown import clean_image_markup, image_reference_labels
from src.pec.knowledge.paths import SOURCES_DIR, SUPPORTED_SUFFIXES, resolve_source_path, source_relative


@dataclass
class ChunkDraft:
    chunk_id: str
    text: str
    source_ref: str
    folder: str
    loc: str
    method: str  # text | table


def _make_chunk_id(source_ref: str, loc: str, text: str) -> str:
    digest = hashlib.sha1(f"{source_ref}|{loc}|{text}".encode()).hexdigest()[:12]
    return f"chunk-{digest}"


def _folder_name(source_ref: str) -> str:
    parts = source_ref.replace("\\", "/").split("/")
    return parts[0] if len(parts) > 1 else ""


def _serialize_table(headers: list[str], row: list[str]) -> str:
    headers = [clean_image_markup(value) for value in headers]
    row = [clean_image_markup(value) for value in row]
    pairs = []
    for i, value in enumerate(row):
        if not (value or "").strip():
            continue
        key = headers[i] if i < len(headers) and headers[i].strip() else f"col_{i+1}"
        pairs.append(f"{key}: {value.strip()}")
    return " | ".join(pairs)


def _walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


def _extract_pptx_chunks(path: Path, source_ref: str, access_token: str) -> list[ChunkDraft]:
    prs = Presentation(str(path))
    chunks: list[ChunkDraft] = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        loc = f"Slide {slide_no}"
        texts: list[str] = []
        table_no = 0
        slide_tables = []
        for shape in _walk_shapes(slide.shapes):
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                table_no += 1
                table = shape.table
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                header_text = " | ".join(clean_image_markup(value) for value in headers)
                if header_text.strip(" |"):
                    header_loc = f"{loc} Table {table_no} row 1"
                    slide_tables.append(ChunkDraft(_make_chunk_id(source_ref, header_loc, header_text),
                        header_text, source_ref, _folder_name(source_ref), header_loc, "table"))
                for row_idx in range(1, len(table.rows)):
                    row = table.rows[row_idx]
                    cells = [cell.text.strip() for cell in row.cells]
                    row_text = _serialize_table(headers, cells)
                    if row_text:
                        row_loc = f"{loc} Table {table_no} row {row_idx + 1}"
                        slide_tables.append(
                            ChunkDraft(
                                chunk_id=_make_chunk_id(source_ref, row_loc, row_text),
                                text=row_text,
                                source_ref=source_ref,
                                folder=_folder_name(source_ref),
                                loc=row_loc,
                                method="table",
                            )
                        )

        body = "\n".join(texts).strip()
        method = "text"
        if body:
            chunks.append(
                ChunkDraft(
                    chunk_id=_make_chunk_id(source_ref, loc, body),
                    text=body,
                    source_ref=source_ref,
                    folder=_folder_name(source_ref),
                    loc=loc,
                    method=method,
                )
            )
        chunks.extend(slide_tables)
    return chunks


def _extract_docx_chunks(path: Path, source_ref: str) -> list[ChunkDraft]:
    doc = Document(str(path))
    chunks: list[ChunkDraft] = []

    # Iterate XML children to preserve paragraphs interleaved with tables.
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    para_no = table_idx = 0
    for element in doc.element.body:
        if element.tag.endswith('}tbl'):
            table_idx += 1
            table = Table(element, doc)
            if not table.rows:
                continue
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            header_text = " | ".join(clean_image_markup(value) for value in headers)
            if header_text.strip(" |"):
                header_loc = f"Table {table_idx} row 1"
                chunks.append(ChunkDraft(_make_chunk_id(source_ref, header_loc, header_text),
                    header_text, source_ref, _folder_name(source_ref), header_loc, "table"))
            for row_idx, row in enumerate(table.rows[1:], start=2):
                text = _serialize_table(headers, [cell.text.strip() for cell in row.cells])
                loc = f"Table {table_idx} row {row_idx}"
                if text:
                    chunks.append(ChunkDraft(_make_chunk_id(source_ref, loc, text), text,
                        source_ref, _folder_name(source_ref), loc, "table"))
            continue
        if not element.tag.endswith('}p'):
            continue
        para_no += 1
        text = Paragraph(element, doc).text.strip()
        if not text:
            continue
        loc = f"Paragraph {para_no}"
        chunks.append(
            ChunkDraft(
                chunk_id=_make_chunk_id(source_ref, loc + text[:40], text),
                text=text,
                source_ref=source_ref,
                folder=_folder_name(source_ref),
                loc=loc,
                method="text",
            )
        )

    return chunks


def _extract_pdf_chunks(path: Path, source_ref: str, access_token: str) -> list[ChunkDraft]:
    reader = PdfReader(str(path))
    chunks: list[ChunkDraft] = []

    for page_no, page in enumerate(reader.pages, start=1):
        loc = f"Page {page_no}"
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(
                ChunkDraft(
                    chunk_id=_make_chunk_id(source_ref, loc, text),
                    text=text,
                    source_ref=source_ref,
                    folder=_folder_name(source_ref),
                    loc=loc,
                    method="text",
                )
            )
    return chunks


def _extract_xlsx_chunks(path: Path, source_ref: str) -> list[ChunkDraft]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[ChunkDraft] = []
    try:
        for sheet in wb.worksheets:
            headers = None
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [str(c).strip() if c is not None else "" for c in row]
                if not any(cells):
                    continue
                if headers is None:
                    headers = cells
                    text = "Headers: " + " | ".join(cells)
                else:
                    text = _serialize_table(headers, cells)
                loc = f"Sheet {sheet.title} row {row_no}"
                chunks.append(
                    ChunkDraft(
                        chunk_id=_make_chunk_id(source_ref, loc, text),
                        text=text,
                        source_ref=source_ref,
                        folder=_folder_name(source_ref),
                        loc=loc,
                        method="table",
                    )
                )
    finally:
        wb.close()
    return chunks


def _ingest_file(path: Path, *, access_token: str | None = None) -> list[ChunkDraft]:
    source_ref = source_relative(path)
    resolved = resolve_source_path(source_ref)
    token = None  # Compatibility argument only: parsing never invokes a model.
    suffix = resolved.suffix.lower()

    if suffix == ".pptx":
        return _extract_pptx_chunks(resolved, source_ref, token)
    if suffix == ".docx":
        return _extract_docx_chunks(resolved, source_ref)
    if suffix == ".pdf":
        return _extract_pdf_chunks(resolved, source_ref, token)
    if suffix == ".xlsx":
        return _extract_xlsx_chunks(resolved, source_ref)
    return []


def ingest_file(path: Path, *, access_token: str | None = None) -> list[ChunkDraft]:
    result = []
    drafts = _ingest_file(path, access_token=access_token)
    image_refs = image_reference_labels("\n".join(c.text for c in drafts))
    for chunk in drafts:
        cleaned = clean_image_markup(chunk.text, image_refs=image_refs)
        if not cleaned:
            continue
        if cleaned != chunk.text:
            chunk.text = cleaned
            chunk.chunk_id = _make_chunk_id(chunk.source_ref, chunk.loc, cleaned)
        if len(chunk.text) <= 6000:
            result.append(chunk)
            continue
        for offset in range(0, len(chunk.text), 5600):
            text = chunk.text[offset:offset + 6000]
            loc = f"{chunk.loc} Part {offset // 5600 + 1}"
            result.append(ChunkDraft(_make_chunk_id(chunk.source_ref, loc, text), text,
                                     chunk.source_ref, chunk.folder, loc, chunk.method))
    return result


def scan_source_files(scope: str = "") -> list[Path]:
    from src.pec.knowledge.sources import LocalSourceProvider
    return [document.path for document in LocalSourceProvider().discover(scope)]


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def chunk_draft_to_dict(chunk: ChunkDraft) -> dict:
    return asdict(chunk)
