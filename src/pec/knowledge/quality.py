"""Offline triage; flags describe incomplete coverage, not successful OCR."""
from pathlib import Path

from src.pec.knowledge.chunk_ingest import _walk_shapes


def inspect_quality(path: Path, chunks: list[dict]) -> list[dict]:
    warnings = []
    if not chunks:
        warnings.append({"loc": "文件", "reason": "无可提取文本，已跳过；不进行 OCR 或视觉识别"})
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        for number, slide in enumerate(Presentation(str(path)).slides, 1):
            shapes = list(_walk_shapes(slide.shapes))
            reasons = []
            if any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes):
                reasons.append("包含图片，图片内文字未提取")
            if any(getattr(s, "has_chart", False) or s.shape_type in
                   {MSO_SHAPE_TYPE.DIAGRAM, MSO_SHAPE_TYPE.LINE} for s in shapes):
                reasons.append("包含图表或连接关系，仅保留可提取文字")
            if sum(s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for s in shapes) >= 3:
                reasons.append("存在多形状布局，仅保留文字，不推断流程或层级关系")
            if not any(c["loc"].startswith(f"Slide {number} ") or c["loc"] == f"Slide {number}" for c in chunks):
                reasons.append("本页未提取到文字，已跳过")
            if reasons:
                warnings.append({"loc": f"Slide {number}", "reason": "；".join(reasons)})
    elif path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(path))
        if doc.element.xpath(".//w:drawing | .//w:pict"):
            warnings.append({"loc": "Document", "reason": "包含图片或绘图，已跳过视觉内容"})
    elif path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        import fitz
        with fitz.open(str(path)) as pdf:
            drawing_pages = {i + 1 for i, page in enumerate(pdf) if page.get_drawings()}
        for number, page in enumerate(PdfReader(str(path)).pages, 1):
            text = (page.extract_text() or "").strip()
            if len(text) < 80 or bool(page.images) or number in drawing_pages:
                warnings.append({"loc": f"Page {number}", "reason": "文字稀少或包含图片/绘图，仅保留可提取文字，不进行 OCR"})
    return warnings
