"""按 source_ref + loc 读取 sources 中的原文片段（比 chunk 更完整）。"""

from __future__ import annotations

import re
from pathlib import Path

from src.pec.knowledge.paths import SUPPORTED_SUFFIXES, resolve_source_path


def normalize_excerpt_loc(loc: str) -> str:
    """将 chunk loc 规范为可读整页/整表的位置（如 Slide 6 Table row 3 → Slide 6）。"""
    loc = (loc or "").strip()
    if m := re.match(r"(Slide\s+\d+)", loc, re.I):
        return m.group(1)
    if m := re.match(r"(Page\s+\d+)", loc, re.I):
        return m.group(1)
    if m := re.match(r"(Table\s+\d+)", loc, re.I):
        return m.group(1)
    return loc


def read_source_excerpt(source_ref: str, loc: str) -> dict:
    """返回该位置的完整页/幻灯片/段落级原文（供 UI 展示）。"""
    path = resolve_source_path(source_ref)
    suffix = path.suffix.lower()
    loc = normalize_excerpt_loc(loc)

    if suffix == ".pptx":
        return _read_pptx_loc(path, source_ref, loc)
    if suffix == ".docx":
        return _read_docx_loc(path, source_ref, loc)
    if suffix == ".pdf":
        return _read_pdf_loc(path, source_ref, loc)
    if suffix == ".xlsx":
        return {"source_ref": source_ref, "loc": loc, "text": "[xlsx 暂不支持按 loc 展开全文]"}
    raise ValueError(f"不支持的文件类型: {suffix}")


def _read_pptx_loc(path: Path, source_ref: str, loc: str) -> dict:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("读取 PPTX 原文需要安装 python-pptx。") from exc
    m = re.match(r"Slide\s+(\d+)", loc, re.I)
    if not m:
        return {"source_ref": source_ref, "loc": loc, "text": _read_pptx_all(path)}
    slide_no = int(m.group(1))
    prs = Presentation(str(path))
    if slide_no < 1 or slide_no > len(prs.slides):
        raise ValueError(f"幻灯片不存在: {loc}")
    slide = prs.slides[slide_no - 1]
    parts: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            parts.append(shape.text.strip())
    text = "\n\n".join(parts).strip() or "（该页无提取到文字，可能为纯图片页）"
    return {"source_ref": source_ref, "loc": f"Slide {slide_no}", "text": text}


def _read_pptx_all(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("读取 PPTX 原文需要安装 python-pptx。") from exc
    prs = Presentation(str(path))
    blocks = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text.strip())
        if lines:
            blocks.append(f"--- Slide {i} ---\n" + "\n".join(lines))
    return "\n\n".join(blocks)


def _read_docx_loc(path: Path, source_ref: str, loc: str) -> dict:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("读取 DOCX 原文需要安装 python-docx。") from exc
    if loc.startswith("Table"):
        doc = Document(str(path))
        m = re.match(r"Table\s+(\d+)", loc, re.I)
        if not m:
            raise ValueError(f"无法解析 loc: {loc}")
        idx = int(m.group(1)) - 1
        if idx < 0 or idx >= len(doc.tables):
            raise ValueError(f"表格不存在: {loc}")
        table = doc.tables[idx]
        rows = []
        for row in table.rows:
            rows.append(" | ".join(c.text.strip() for c in row.cells))
        return {"source_ref": source_ref, "loc": loc, "text": "\n".join(rows)}
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return {"source_ref": source_ref, "loc": loc or "全文", "text": "\n\n".join(paras)}


def _read_pdf_loc(path: Path, source_ref: str, loc: str) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("读取 PDF 原文需要安装 pypdf。") from exc
    m = re.match(r"Page\s+(\d+)", loc, re.I)
    reader = PdfReader(str(path))
    if m:
        page_no = int(m.group(1))
        if page_no < 1 or page_no > len(reader.pages):
            raise ValueError(f"页码不存在: {loc}")
        text = (reader.pages[page_no - 1].extract_text() or "").strip()
        return {"source_ref": source_ref, "loc": f"Page {page_no}", "text": text or "（该页无文本）"}
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        t = (page.extract_text() or "").strip()
        if t:
            parts.append(f"--- Page {i} ---\n{t}")
    return {"source_ref": source_ref, "loc": "全文", "text": "\n\n".join(parts)}
