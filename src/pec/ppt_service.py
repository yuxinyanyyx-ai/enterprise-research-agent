"""
PEC PPT解析业务层
"""

from pathlib import Path
from pptx import Presentation


def parse_pptx_structure_service(
    file_path: str,
) -> str:
    """
    解析PPT结构，提取每页文本内容。
    """

    path = Path(file_path)

    if not path.exists():
        return f"PPT文件不存在: {file_path}"

    if path.suffix.lower() != ".pptx":
        return "目前只支持 .pptx 文件"

    try:
        prs = Presentation(path)

        result = []

        for idx, slide in enumerate(prs.slides, start=1):

            texts = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        texts.append(text)

            slide_text = "\n".join(texts)

            result.append(
                f"""
===== Slide {idx} =====

{slide_text if slide_text else "(无文本)"}
"""
            )

        return "\n".join(result).strip()


    except Exception as exc:

        return (
            f"PPT解析失败: {exc}"
        )