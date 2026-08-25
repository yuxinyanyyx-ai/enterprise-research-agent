"""Service boundary for parsing documents and running confirmed DMF queries."""

from __future__ import annotations

import shutil
from collections.abc import Callable
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from src.Apollo.workflow_client import extract_dmf_query_params
from src.dmf_query.multi_query_service import search_dmf_queries
from src.mineru.services.mineru_client import BatchParseResult, MinerUClient
from src.schemas.document_dmf import (
	DocumentArtifact,
	ExtractedDMFQueryBatch,
)
from src.settings import Settings, get_settings


class DocumentDMFError(RuntimeError):
	"""Document workflow input or processing failed."""


class DocumentDMFService:
	"""Coordinate MinerU, Apollo Studio Flow, and confirmed DMF searches."""

	def __init__(
		self,
		settings: Settings | None = None,
		*,
		mineru_factory: Callable[[Settings], MinerUClient] = MinerUClient,
		extractor: Callable[[str], dict[str, Any]] = extract_dmf_query_params,
		searcher: Callable[..., dict[str, Any]] = search_dmf_queries,
	) -> None:
		self.settings = settings or get_settings()
		self._mineru_factory = mineru_factory
		self._extractor = extractor
		self._searcher = searcher

	def parse_document(self, file_path: str | Path) -> DocumentArtifact:
		path = Path(file_path).expanduser().resolve()
		if not path.is_file():
			raise DocumentDMFError(f"文件不存在：{path}")
		if path.stat().st_size == 0:
			raise DocumentDMFError(f"文件为空：{path.name}")
		if path.stat().st_size > self.settings.max_file_size_bytes:
			raise DocumentDMFError(f"文件超过大小限制：{path.name}")

		suffix = path.suffix.lower()
		if suffix in {".md", ".txt"}:
			return self._store_text_document(path)
		if suffix not in self.settings.allowed_extensions:
			raise DocumentDMFError(f"不支持的文件格式：{path.name}")
		if suffix in {".xlsx", ".xls"}:
			return self._store_excel_document(path)

		with self._mineru_factory(self.settings) as client:
			result = client.parse_files([path])
		return self._artifact_from_mineru(path, result)

	def extract_query(
		self,
		artifact: DocumentArtifact | dict[str, Any],
	) -> ExtractedDMFQueryBatch:
		validated_artifact = DocumentArtifact.model_validate(artifact)
		markdown_path = self._trusted_markdown_path(validated_artifact.markdown_path)
		try:
			markdown = markdown_path.read_text(encoding="utf-8-sig")
		except OSError as exc:
			raise DocumentDMFError(f"读取 Markdown 失败：{exc}") from exc
		if not markdown.strip():
			raise DocumentDMFError("Markdown 内容为空")

		try:
			return ExtractedDMFQueryBatch.model_validate(self._extractor(markdown))
		except Exception as exc:
			raise DocumentDMFError(f"Flow 提取 DMF 条件失败：{exc}") from exc

	def execute_confirmed_query(
		self,
		query: ExtractedDMFQueryBatch,
	) -> dict[str, Any]:
		validated = ExtractedDMFQueryBatch.model_validate(query)
		results = [
			self._searcher(
				dmf_no=item.dmf_no,
				applicant_name=item.applicant_name,
				ingredients=item.ingredients,
			)
			for item in validated.queries
		]
		return self._merge_query_results(results)

	@staticmethod
	def _merge_query_results(results: list[dict[str, Any]]) -> dict[str, Any]:
		if len(results) == 1:
			return results[0]

		query_count = 0
		success_count = 0
		failed_count = 0
		merged_results: list[dict[str, Any]] = []
		for result in results:
			query = result.get("query") or {}
			ingredients = query.get("ingredients") or [""]
			attempt_count = max(
				int(result.get("query_count", 0)),
				len(ingredients),
				1,
			)
			query_count += attempt_count
			success_count += int(result.get("success_count", 0))
			result_failed_count = int(result.get("failed_count", 0))
			if not result.get("success") and result_failed_count == 0:
				result_failed_count = attempt_count
			failed_count += result_failed_count

			nested_results = result.get("results") or []
			if nested_results:
				merged_results.extend(nested_results)
			elif not result.get("success"):
				merged_results.extend(
					{
						"success": False,
						"message": result.get("message", "查询失败"),
						"query": {
							"dmf_no": query.get("dmf_no", ""),
							"applicant_name": query.get("applicant_name", ""),
							"ingredient": ingredient,
						},
						"total": 0,
						"total_pages": 0,
						"records": [],
					}
					for ingredient in ingredients
				)
		return {
			"success": failed_count == 0,
			"message": f"批量查询完成：成功 {success_count}，失败 {failed_count}",
			"query_count": query_count,
			"success_count": success_count,
			"failed_count": failed_count,
			"total_records": sum(
				int(result.get("total_records", 0)) for result in results
			),
			"results": merged_results,
		}

	def _store_text_document(self, path: Path) -> DocumentArtifact:
		try:
			text = path.read_text(encoding="utf-8-sig")
		except (OSError, UnicodeError) as exc:
			raise DocumentDMFError(f"读取文档失败：{exc}") from exc
		return self._store_markdown_document(path, text)

	def _store_excel_document(self, path: Path) -> DocumentArtifact:
		try:
			sheets = (
				self._read_xlsx_sheets(path)
				if path.suffix.lower() == ".xlsx"
				else self._read_xls_sheets(path)
			)
			markdown = self._sheets_to_markdown(sheets)
		except DocumentDMFError:
			raise
		except Exception as exc:
			raise DocumentDMFError(f"解析 Excel 文件失败：{path.name}：{exc}") from exc
		return self._store_markdown_document(path, markdown)

	@staticmethod
	def _read_xlsx_sheets(path: Path) -> list[tuple[str, list[list[Any]]]]:
		workbook = load_workbook(
			path,
			read_only=True,
			data_only=True,
			keep_links=False,
		)
		try:
			return [
				(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
				for sheet in workbook.worksheets
			]
		finally:
			workbook.close()

	@staticmethod
	def _read_xls_sheets(path: Path) -> list[tuple[str, list[list[Any]]]]:
		try:
			import xlrd
		except ImportError as exc:
			raise DocumentDMFError(
				"读取 .xls 需要安装 xlrd，请执行 pip install -r requirements.txt"
			) from exc

		workbook = xlrd.open_workbook(path, on_demand=True)
		try:
			sheets: list[tuple[str, list[list[Any]]]] = []
			for sheet in workbook.sheets():
				rows: list[list[Any]] = []
				for row_index in range(sheet.nrows):
					row: list[Any] = []
					for column_index in range(sheet.ncols):
						cell = sheet.cell(row_index, column_index)
						value = cell.value
						if cell.ctype == xlrd.XL_CELL_DATE:
							value = xlrd.xldate_as_datetime(value, workbook.datemode)
						row.append(value)
					rows.append(row)
				sheets.append((sheet.name, rows))
			return sheets
		finally:
			workbook.release_resources()

	@classmethod
	def _sheets_to_markdown(
		cls,
		sheets: list[tuple[str, list[list[Any]]]],
	) -> str:
		parts: list[str] = []
		for sheet_name, raw_rows in sheets:
			rows = [cls._trim_excel_row(row) for row in raw_rows]
			rows = [row for row in rows if row]
			if not rows:
				continue

			width = max(len(row) for row in rows)
			headers = cls._normalize_headers(rows[0], width)
			body = [row + [""] * (width - len(row)) for row in rows[1:]]
			title = cls._markdown_cell(sheet_name) or "Sheet"
			lines = [
				f"## {title}",
				"",
				"| " + " | ".join(headers) + " |",
				"| " + " | ".join("---" for _ in range(width)) + " |",
			]
			lines.extend(
				"| " + " | ".join(cls._markdown_cell(value) for value in row) + " |"
				for row in body
			)
			parts.append("\n".join(lines))

		if not parts:
			raise DocumentDMFError("Excel 工作簿中没有可读取的单元格数据")
		return "\n\n".join(parts)

	@classmethod
	def _trim_excel_row(cls, row: list[Any]) -> list[Any]:
		values = list(row)
		while values and not cls._cell_has_value(values[-1]):
			values.pop()
		return values if any(cls._cell_has_value(value) for value in values) else []

	@staticmethod
	def _cell_has_value(value: Any) -> bool:
		return value is not None and (not isinstance(value, str) or bool(value.strip()))

	@classmethod
	def _normalize_headers(cls, row: list[Any], width: int) -> list[str]:
		seen: dict[str, int] = {}
		headers: list[str] = []
		for index in range(width):
			base = cls._markdown_cell(row[index] if index < len(row) else "")
			base = base or f"Column {index + 1}"
			key = base.casefold()
			seen[key] = seen.get(key, 0) + 1
			headers.append(base if seen[key] == 1 else f"{base} {seen[key]}")
		return headers

	@staticmethod
	def _markdown_cell(value: Any) -> str:
		if value is None:
			return ""
		if isinstance(value, (datetime, date, time)):
			text = value.isoformat()
		elif isinstance(value, float) and value.is_integer():
			text = str(int(value))
		else:
			text = str(value).strip()
		return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>")

	def _store_markdown_document(self, path: Path, text: str) -> DocumentArtifact:
		if not text.strip():
			raise DocumentDMFError("文档内容为空")
		if len(text.encode("utf-8")) > self.settings.max_markdown_size_bytes:
			raise DocumentDMFError("转换后的 Markdown 超过大小限制")

		document_id = uuid4().hex
		document_dir = self.settings.result_dir / "documents" / document_id
		markdown_path = document_dir / "document.md"
		try:
			document_dir.mkdir(parents=True, exist_ok=False)
			markdown_path.write_text(text, encoding="utf-8")
		except OSError as exc:
			shutil.rmtree(document_dir, ignore_errors=True)
			raise DocumentDMFError(f"保存 Markdown 失败：{exc}") from exc
		return DocumentArtifact(
			document_id=document_id,
			file_name=path.name,
			source_path=str(path),
			markdown_path=str(markdown_path.resolve()),
		)

	def _artifact_from_mineru(
		self,
		path: Path,
		result: BatchParseResult,
	) -> DocumentArtifact:
		if not result.succeeded:
			error = result.files[0].error if result.files else "MinerU 未返回文件结果"
			raise DocumentDMFError(error or "MinerU 文档解析失败")
		parsed = result.succeeded[0]
		if parsed.markdown_path is None:
			raise DocumentDMFError("MinerU 未生成 Markdown 文件")
		markdown_path = parsed.markdown_path.resolve()
		self._trusted_markdown_path(markdown_path)
		return DocumentArtifact(
			document_id=parsed.data_id,
			file_name=path.name,
			source_path=str(path),
			markdown_path=str(markdown_path),
		)

	def _trusted_markdown_path(self, value: str | Path) -> Path:
		path = Path(value).resolve()
		root = self.settings.result_dir.resolve()
		if not path.is_relative_to(root):
			raise DocumentDMFError("拒绝读取结果目录之外的 Markdown")
		if not path.is_file():
			raise DocumentDMFError(f"Markdown 文件不存在：{path}")
		return path