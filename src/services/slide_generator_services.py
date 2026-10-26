"""Generate and modify project slide 1 from uploaded data slides (LPT + milestone)."""

from __future__ import annotations

import json
import os
import re
import uuid
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from src.llm.apollo import create_apollo_llm, invoke_with_gateway_retries
from src.settings import load_settings
from src.tools.project_slide_tool.fill import run_fill
from src.tools.project_slide_tool.lpt import extract_bottom_tables
from src.tools.project_slide_tool.milestone import extract_slide_title, parse_milestone_slide
from src.tools.project_slide_tool.models import BottomTables, FillConfig, FillResult, TrackData

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
DEFAULT_TEMPLATE = ASSETS_DIR / "project_slide_template.pptx"


def template_path() -> Path:
    env = os.getenv("PROJECT_SLIDE_TEMPLATE", "").strip()
    if env:
        path = Path(env)
        if path.is_file():
            return path
        raise FileNotFoundError(f"PROJECT_SLIDE_TEMPLATE 不存在: {path}")
    if DEFAULT_TEMPLATE.is_file():
        return DEFAULT_TEMPLATE
    raise FileNotFoundError(
        f"未找到 project slide 模板: {DEFAULT_TEMPLATE}。"
        "请将 project_slide_template.pptx 放入 src/services/assets/。"
    )


def _validate_identifier(value: str, label: str) -> str:
    if not re.fullmatch(r"[0-9a-f]{32}", value or ""):
        raise ValueError(f"{label} 无效或已失效")
    return value


def _session_directory(base: Path, session_id: str) -> Path:
    session_id = _validate_identifier(session_id, "Agent 会话")
    path = (base / "project_slide" / session_id).resolve()
    root = (base / "project_slide").resolve()
    if path.parent != root:
        raise ValueError("Agent 会话路径无效")
    path.mkdir(parents=True, exist_ok=True)
    return path


def _upload_directory(session_id: str) -> Path:
    return _session_directory(load_settings(require_token=False).upload_dir, session_id)


def _output_directory(session_id: str) -> Path:
    return _session_directory(load_settings(require_token=False).result_dir, session_id)


def store_uploaded_pptx(session_id: str, filename: str, content: bytes) -> dict[str, str]:
    if Path(filename).suffix.lower() != ".pptx":
        raise ValueError("仅支持 .pptx 文件")
    if not content:
        raise ValueError("PPTX 文件为空")
    try:
        Presentation(BytesIO(content))
    except Exception as exc:
        raise ValueError("上传文件不是有效的 PPTX") from exc

    file_id = uuid.uuid4().hex
    directory = _upload_directory(session_id)
    path = directory / f"{file_id}.pptx"
    path.write_bytes(content)
    metadata = {"file_id": file_id, "filename": Path(filename).name}
    path.with_suffix(".json").write_text(
        json.dumps(metadata, ensure_ascii=False), encoding="utf-8"
    )
    return metadata


def uploaded_file_path(session_id: str, file_id: str) -> Path:
    file_id = _validate_identifier(file_id, "文件 ID")
    path = _upload_directory(session_id) / f"{file_id}.pptx"
    if not path.is_file():
        raise FileNotFoundError("当前 Agent 会话中找不到该 PPTX")
    return path


def uploaded_pptx_path(session_id: str, file_id: str) -> Path:
    return uploaded_file_path(session_id, file_id)


def uploaded_filename(session_id: str, file_id: str) -> str:
    path = uploaded_file_path(session_id, file_id).with_suffix(".json")
    metadata = json.loads(path.read_text(encoding="utf-8"))
    return str(metadata.get("filename") or "uploaded.pptx")


def job_output_path(session_id: str, job_id: str) -> Path:
    job_id = _validate_identifier(job_id, "Project Slide job ID")
    return _output_directory(session_id) / f"{job_id}.pptx"


def load_job_meta(session_id: str, job_id: str) -> dict:
    job_id = _validate_identifier(job_id, "Project Slide job ID")
    path = job_output_path(session_id, job_id).with_suffix(".json")
    if not path.is_file():
        raise FileNotFoundError("当前 Agent 会话中找不到该 Project Slide")
    metadata = json.loads(path.read_text(encoding="utf-8"))
    if metadata.get("job_id") != job_id or metadata.get("session_id") != session_id:
        raise ValueError("该 Project Slide 不属于当前 Agent 会话")
    return metadata


def save_job_meta(session_id: str, job_id: str, metadata: dict) -> None:
    output_path = job_output_path(session_id, job_id)
    if not output_path.is_file():
        raise FileNotFoundError("Project Slide 输出文件不存在")
    metadata["session_id"] = session_id
    metadata["job_id"] = job_id
    metadata["filename"] = f"project_slide_{job_id[:8]}.pptx"
    metadata["download_url"] = (
        f"/api/agent/sessions/{session_id}/project-slides/{job_id}"
    )
    meta_path = output_path.with_suffix(".json")
    temp_path = meta_path.with_suffix(".json.tmp")
    temp_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
    temp_path.replace(meta_path)


def resolve_data_slides(pptx_path: Path) -> tuple[int, int]:
    """Map uploaded pages to LPT / milestone slide indices (1-based)."""
    prs = Presentation(str(pptx_path))
    count = len(prs.slides)
    if count < 2:
        raise ValueError("上传的 PPTX 至少需要 2 页（Status/Risk 页 + Milestone 页）")
    if count == 2:
        return 1, 2
    return 2, 3


def load_slide_data(
    pptx_path: Path,
    lpt_slide: int,
    milestone_slide: int,
) -> tuple[TrackData, TrackData, str, BottomTables, str]:
    data_prs = Presentation(str(pptx_path))
    if len(data_prs.slides) < max(lpt_slide, milestone_slide):
        raise ValueError(f"数据文件至少需要 {max(lpt_slide, milestone_slide)} 页")

    upper, lower = parse_milestone_slide(data_prs.slides[milestone_slide - 1])
    slide_title = extract_slide_title(data_prs.slides[milestone_slide - 1])
    bottom, bottom_source = extract_bottom_tables(pptx_path, lpt_slide - 1)
    return upper, lower, slide_title, bottom, bottom_source


def generate_project_slide(
    data_path: Path,
    output_path: Path | None = None,
    *,
    lpt_slide: int | None = None,
    milestone_slide: int | None = None,
    slide_title: str | None = None,
    upper: TrackData | None = None,
    lower: TrackData | None = None,
    bottom: BottomTables | None = None,
) -> FillResult:
    tpl = template_path()
    if lpt_slide is None or milestone_slide is None:
        lpt_slide, milestone_slide = resolve_data_slides(data_path)
    out = output_path or (data_path.parent / "project_slide.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    return run_fill(
        FillConfig(
            template=str(tpl),
            data=str(data_path),
            output=str(out),
            data_lpt_slide=lpt_slide,
            data_milestone_slide=milestone_slide,
            slide_title=slide_title,
            upper=upper,
            lower=lower,
            bottom=bottom,
        )
    )


def _slide_data_from_meta(meta: dict) -> tuple[TrackData, TrackData, str, BottomTables]:
    if meta.get("upper") and meta.get("lower") and meta.get("bottom"):
        return (
            TrackData.from_dict(meta["upper"]),
            TrackData.from_dict(meta["lower"]),
            str(meta.get("slide_title") or ""),
            BottomTables.from_dict(meta["bottom"]),
        )

    pptx_path = uploaded_pptx_path(meta["session_id"], meta["file_id"])
    upper, lower, slide_title, bottom, _ = load_slide_data(
        pptx_path,
        int(meta["lpt_slide"]),
        int(meta["milestone_slide"]),
    )
    return upper, lower, slide_title, bottom


def _format_track_markdown(title: str, track: TrackData) -> str:
    lines = [f"### {title}", f"Study label: {track.study_label or '—'}", ""]
    lines.append("| Milestone | Date | Note |")
    lines.append("|---|---|---|")
    for ms in track.event_order:
        ev = track.events.get(ms)
        if not ev:
            continue
        lines.append(f"| {ms} | {ev.date or ''} | {ev.note or ''} |")
    return "\n".join(lines)


def _format_bottom_markdown(bottom: BottomTables) -> str:
    lines = [
        "### Bottom Tables",
        "",
        "| | Status | Next Steps |",
        "|---|---|---|",
        f"| **Global** | {bottom.global_status or ''} | {bottom.global_next or ''} |",
        f"| **China** | {bottom.china_status or ''} | {bottom.china_next or ''} |",
        "",
        "| Risk / Issue | Mitigation / Action |",
        "|---|---|",
    ]
    if bottom.risks:
        for left, right in bottom.risks:
            lines.append(f"| {left} | {right} |")
    else:
        lines.append("| | |")
    return "\n".join(lines)


def _extract_json_object(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("LLM 未返回有效 JSON")
    return json.loads(text[start : end + 1])


def _setup_apollo_llm():
    llm = create_apollo_llm()
    llm.streaming = False
    llm.temperature = 0
    return llm


def apply_slide_feedback(
    slide_title: str,
    upper: TrackData,
    lower: TrackData,
    bottom: BottomTables,
    feedback: str,
) -> tuple[str, TrackData, TrackData, BottomTables]:
    feedback = (feedback or "").strip()
    if not feedback:
        raise ValueError("修改意见不能为空")

    prompt = f"""You are editing an existing Project Slide. Apply the user's feedback to the current content.

Rules:
- Keep the same milestone names and overall structure. Do NOT add or remove milestones.
- Only change text fields: slide title, study labels, milestone dates, milestone notes, bottom table cells.
- If the user asks to change a date, update the date text. Layout will follow the date automatically.
- Preserve bullet formatting in bottom tables using `• item` and `<br>` between bullets when multiple bullets exist.
- Output language for slide content: English.
- Respond with ONLY valid JSON, no markdown fences.

Current content:

## Slide Title
{slide_title or "—"}

{_format_track_markdown("Timeline Upper Track", upper)}

{_format_track_markdown("Timeline Lower Track", lower)}

{_format_bottom_markdown(bottom)}

User feedback:
{feedback}

JSON schema:
{{
  "slide_title": "...",
  "upper": {{"study_label": "...", "event_order": ["..."], "events": {{"FPI": {{"date": "...", "note": "..."}}}}}},
  "lower": {{"study_label": "...", "event_order": ["..."], "events": {{}}}},
  "bottom": {{
    "global_status": "...",
    "global_next": "...",
    "china_status": "...",
    "china_next": "...",
    "risks": [["risk", "mitigation"]]
  }}
}}"""

    llm = _setup_apollo_llm()
    response = invoke_with_gateway_retries(
        llm, prompt, retries=2,
        operation_name="llm.slide_generation", component="slide_generator",
    )
    content = response.content if hasattr(response, "content") else str(response)
    if isinstance(content, list):
        content = "\n".join(
            p.get("text", "") if isinstance(p, dict) else str(p) for p in content
        )
    payload = _extract_json_object(str(content))

    new_upper = TrackData.from_dict(payload.get("upper"))
    new_lower = TrackData.from_dict(payload.get("lower"))
    new_bottom = BottomTables.from_dict(payload.get("bottom"))
    new_title = str(payload.get("slide_title") or slide_title)

    if not new_upper.events:
        new_upper = upper
    if not new_lower.events:
        new_lower = lower
    if new_bottom.is_empty():
        new_bottom = bottom

    return new_title, new_upper, new_lower, new_bottom


def _build_job_payload(
    *,
    session_id: str,
    job_id: str,
    file_id: str,
    original_filename: str,
    lpt_slide: int,
    milestone_slide: int,
    slide_title: str,
    upper: TrackData,
    lower: TrackData,
    bottom: BottomTables,
    result: FillResult,
    parent_job_id: str | None = None,
) -> dict:
    payload = {
        "ok": True,
        "session_id": session_id,
        "job_id": job_id,
        "file_id": file_id,
        "filename": "output.pptx",
        "source_filename": original_filename,
        "download_url": f"/api/agent/sessions/{session_id}/project-slides/{job_id}",
        "lpt_slide": lpt_slide,
        "milestone_slide": milestone_slide,
        "slide_title": slide_title,
        "upper": upper.to_dict(),
        "lower": lower.to_dict(),
        "bottom": bottom.to_dict(),
        "upper_milestones": result.upper_milestones,
        "lower_milestones": result.lower_milestones,
        "bottom_source": result.bottom_source,
        "bottom_risks_count": result.bottom_risks_count,
        "skipped_milestones": result.skipped_milestones,
    }
    if parent_job_id:
        payload["parent_job_id"] = parent_job_id
    return payload


def generate_for_file_id(session_id: str, file_id: str) -> dict:
    """Generate slide 1 from a chat-uploaded pptx (file_id) and register download metadata."""
    pptx_path = uploaded_pptx_path(session_id, file_id)
    original = uploaded_filename(session_id, file_id)
    lpt_slide, milestone_slide = resolve_data_slides(pptx_path)
    upper, lower, slide_title, bottom, _ = load_slide_data(pptx_path, lpt_slide, milestone_slide)

    job_id = uuid.uuid4().hex
    out_path = job_output_path(session_id, job_id)
    temp_path = out_path.with_name(f".{job_id}.tmp.pptx")
    try:
        result = generate_project_slide(
            pptx_path,
            temp_path,
            lpt_slide=lpt_slide,
            milestone_slide=milestone_slide,
            slide_title=slide_title,
            upper=upper,
            lower=lower,
            bottom=bottom,
        )
        temp_path.replace(out_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    payload = _build_job_payload(
        session_id=session_id,
        job_id=job_id,
        file_id=file_id,
        original_filename=original,
        lpt_slide=lpt_slide,
        milestone_slide=milestone_slide,
        slide_title=slide_title,
        upper=upper,
        lower=lower,
        bottom=bottom,
        result=result,
    )
    save_job_meta(session_id, job_id, payload)
    return payload


def modify_for_job_id(session_id: str, job_id: str, feedback: str) -> dict:
    """Modify a previously generated Project Slide and register a new download link."""
    meta = load_job_meta(session_id, job_id)
    upper, lower, slide_title, bottom = _slide_data_from_meta(meta)
    new_title, new_upper, new_lower, new_bottom = apply_slide_feedback(
        slide_title, upper, lower, bottom, feedback
    )

    pptx_path = uploaded_pptx_path(session_id, meta["file_id"])
    lpt_slide = int(meta["lpt_slide"])
    milestone_slide = int(meta["milestone_slide"])
    new_job_id = uuid.uuid4().hex
    out_path = job_output_path(session_id, new_job_id)
    temp_path = out_path.with_name(f".{new_job_id}.tmp.pptx")
    try:
        result = generate_project_slide(
            pptx_path,
            temp_path,
            lpt_slide=lpt_slide,
            milestone_slide=milestone_slide,
            slide_title=new_title,
            upper=new_upper,
            lower=new_lower,
            bottom=new_bottom,
        )
        temp_path.replace(out_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    payload = _build_job_payload(
        session_id=session_id,
        job_id=new_job_id,
        file_id=meta["file_id"],
        original_filename=meta.get("source_filename") or uploaded_filename(session_id, meta["file_id"]),
        lpt_slide=lpt_slide,
        milestone_slide=milestone_slide,
        slide_title=new_title,
        upper=new_upper,
        lower=new_lower,
        bottom=new_bottom,
        result=result,
        parent_job_id=job_id,
    )
    save_job_meta(session_id, new_job_id, payload)
    return payload


def format_generation_result(payload: dict) -> str:
    """Tool/agent-facing summary with download link."""
    skipped = payload.get("skipped_milestones") or {}
    skipped_lines = []
    for track, items in skipped.items():
        if items:
            skipped_lines.append(f"- {track}: {', '.join(items)}")
    skipped_block = "\n".join(skipped_lines) if skipped_lines else "（无）"
    return f"""# Project Slide 生成成功

源文件：{payload.get('source_filename') or payload.get('filename', '')}
数据页：LPT 第 {payload['lpt_slide']} 页，Milestone 第 {payload['milestone_slide']} 页

## 下载链接
下载链接如下，**必须原样输出给用户，不得修改、省略或替换**：
[📥 下载 Project Slide]({payload['download_url']})

## 填充摘要
- 上轨：{' · '.join(payload.get('upper_milestones') or []) or '—'}
- 下轨：{' · '.join(payload.get('lower_milestones') or []) or '—'}
- Risk 条数：{payload.get('bottom_risks_count', 0)}
- 未放置里程碑：{skipped_block}

⚠️ 重要：请用中文告知用户生成已完成，然后直接输出上方的下载链接。**禁止输出任何 JSON 内容。**
"""


def format_modification_result(payload: dict) -> str:
    """Tool/agent-facing summary for modified slide."""
    skipped = payload.get("skipped_milestones") or {}
    skipped_lines = []
    for track, items in skipped.items():
        if items:
            skipped_lines.append(f"- {track}: {', '.join(items)}")
    skipped_block = "\n".join(skipped_lines) if skipped_lines else "（无）"
    parent = payload.get("parent_job_id") or "—"
    return f"""# Project Slide 修改成功

基于上一版 job_id：{parent}
源文件：{payload.get('source_filename') or payload.get('filename', '')}

## 下载链接
下载链接如下，**必须原样输出给用户，不得修改、省略或替换**：
[📥 下载修改后的 Project Slide]({payload['download_url']})

## 填充摘要
- 上轨：{' · '.join(payload.get('upper_milestones') or []) or '—'}
- 下轨：{' · '.join(payload.get('lower_milestones') or []) or '—'}
- Risk 条数：{payload.get('bottom_risks_count', 0)}
- 未放置里程碑：{skipped_block}

⚠️ 重要：请用中文告知用户修改已完成，然后直接输出上方的下载链接。**禁止输出任何 JSON 内容。**
"""
