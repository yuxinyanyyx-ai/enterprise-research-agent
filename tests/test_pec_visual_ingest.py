from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from src.pec.knowledge import chunk_ingest, paths


def _make_pptx(path: Path) -> None:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "A complex portfolio page"
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "status and timeline"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def test_page_vision_is_called_once_per_slide_and_cached(tmp_path, monkeypatch) -> None:
    sources = tmp_path / "sources"
    index = tmp_path / "index"
    source = sources / "PEC1" / "meeting.pptx"
    source.parent.mkdir(parents=True)
    index.mkdir()
    _make_pptx(source)
    monkeypatch.setattr(paths, "SOURCES_DIR", sources)
    monkeypatch.setattr(paths, "INDEX_DIR", index)
    monkeypatch.setattr(chunk_ingest, "INDEX_DIR", index)
    monkeypatch.setenv("PEC_PPT_VISION_MODE", "all")

    image = tmp_path / "slide.png"
    image.write_bytes(b"png")
    monkeypatch.setattr(
        chunk_ingest,
        "_rendered_slide_image",
        lambda source_ref, slide_no: image,
    )
    calls = []
    monkeypatch.setattr(
        chunk_ingest,
        "describe_image",
        lambda blob, **kwargs: calls.append((blob, kwargs)) or f"visual slide {len(calls)}",
    )

    first = chunk_ingest.ingest_file(source, access_token="token")
    second = chunk_ingest.ingest_file(source, access_token="token")

    assert [chunk.method for chunk in first if chunk.method == "vision_page"] == [
        "vision_page",
        "vision_page",
    ]
    assert len(calls) == 2
    cached_text = [chunk.text for chunk in second if chunk.method == "vision_page"]
    assert len(cached_text) == 2
    assert all(f"visual slide {index}" in text for index, text in enumerate(cached_text, 1))


def test_smart_and_none_modes_control_page_vision(tmp_path, monkeypatch) -> None:
    sources = tmp_path / "sources"
    index = tmp_path / "index"
    source = sources / "PEC1" / "meeting.pptx"
    source.parent.mkdir(parents=True)
    index.mkdir()
    _make_pptx(source)
    monkeypatch.setattr(paths, "SOURCES_DIR", sources)
    monkeypatch.setattr(paths, "INDEX_DIR", index)
    monkeypatch.setattr(chunk_ingest, "INDEX_DIR", index)
    monkeypatch.setattr(chunk_ingest, "_rendered_slide_image", lambda *args: image)

    image = tmp_path / "slide.png"
    image.write_bytes(b"png")
    calls = []
    monkeypatch.setattr(
        chunk_ingest,
        "describe_image",
        lambda blob, **kwargs: calls.append(blob) or "visual",
    )

    monkeypatch.setenv("PEC_PPT_VISION_MODE", "none")
    assert not [chunk for chunk in chunk_ingest.ingest_file(source) if chunk.method == "vision_page"]

    monkeypatch.setenv("PEC_PPT_VISION_MODE", "smart")
    assert [chunk for chunk in chunk_ingest.ingest_file(source) if chunk.method == "vision_page"]
    assert calls