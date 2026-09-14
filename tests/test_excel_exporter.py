import pytest
from openpyxl import load_workbook

from src.export_excel.excel_exporter import export_multi_query_result


@pytest.mark.parametrize("filename", ["C:\\outside.xlsx", "..\\outside.xlsx", "CON.xlsx", "bad:name.xlsx", "trailing."])
def test_export_rejects_windows_unsafe_names(tmp_path, filename):
    from tests.test_react_composition import result_data
    with pytest.raises(ValueError):
        export_multi_query_result(result_data(), output_dir=tmp_path, filename=filename)


def test_export_does_not_overwrite_existing_filename(tmp_path):
    from tests.test_react_composition import result_data
    first = export_multi_query_result(result_data(), output_dir=tmp_path, filename="result.xlsx")
    original = first.read_bytes()
    second = export_multi_query_result(result_data(), output_dir=tmp_path, filename="result.xlsx")
    assert first != second
    assert first.read_bytes() == original


def _result() -> dict:
    return {
        "success": True,
        "message": "全部查询完成",
        "query_count": 1,
        "success_count": 1,
        "failed_count": 0,
        "total_records": 1,
        "results": [
            {
                "success": True,
                "query": {"ingredient": "Ibuprofen"},
                "records": [
                    {
                        "dmf_no": "12345",
                        "applicant_name": "Example Pharma",
                        "ingredient": "Ibuprofen",
                        "valid_date": "2027-01-01",
                    }
                ],
            }
        ],
    }


def test_export_creates_expected_workbook(tmp_path) -> None:
    output_path = export_multi_query_result(
        _result(),
        output_dir=tmp_path,
        filename="result",
    )

    workbook = load_workbook(output_path)
    detail = workbook["DMF查询结果"]
    summary = workbook["查询汇总"]

    assert output_path.name == "result.xlsx"
    assert detail.max_column == 7
    assert [cell.value for cell in detail[1]] == [
        "序號",
        "查詢關鍵詞",
        "DMF編號",
        "申請商名",
        "成分名稱",
        "有效日期",
        "查詢狀態",
    ]
    assert [cell.value for cell in detail[2]] == [
        1,
        "Ibuprofen",
        "12345",
        "Example Pharma",
        "Ibuprofen",
        "2027-01-01",
        "查询成功",
    ]
    assert summary[2][1].value == "全部查询完成"


def test_export_rejects_filename_with_directory(tmp_path) -> None:
    try:
        export_multi_query_result(
            _result(),
            output_dir=tmp_path,
            filename="../outside.xlsx",
        )
    except ValueError as exc:
        assert "不能包含目录" in str(exc)
    else:
        raise AssertionError("Expected unsafe filename to be rejected")