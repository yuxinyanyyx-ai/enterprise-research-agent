from __future__ import annotations

import os
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Inches

from src.tools.project_slide_tool import lpt


def _data_pptx() -> bytes:
    presentation = Presentation()
    status_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    status_slide.shapes.add_textbox(0, 0, Inches(9), Inches(4)).text = (
        "Acceptance Test Project. Global Status: Study approved. "
        "Global Next Steps: Begin recruitment. China Status: Sites selected. "
        "China Next Steps: Train sites. Risk: Slow recruitment. "
        "Mitigation: Monitor recruitment weekly."
    )
    milestone_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    milestone_slide.shapes.title.text = "Acceptance Test Project"
    table = milestone_slide.shapes.add_table(
        6, 9, 0, Inches(1), Inches(10), Inches(4)
    ).table
    table.cell(2, 3).text = "Global Test Study"
    table.cell(2, 7).text = "China Test Study"
    for row_index, name, global_date, china_date in (
        (3, "FPI", "Jan 2026", "Mar 2026"),
        (4, "LPI", "Dec 2026", "Feb 2027"),
        (5, "NDA sub.", "Jan 2028", "Mar 2028"),
    ):
        table.cell(row_index, 0).text = name
        table.cell(row_index, 3).text = global_date
        table.cell(row_index, 7).text = china_date
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _slide_text(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    texts = []
    pending = list(presentation.slides[0].shapes)
    while pending:
        shape = pending.pop(0)
        if hasattr(shape, "shapes"):
            pending.extend(shape.shapes)
        if shape.has_text_frame:
            texts.append(shape.text)
        if shape.has_table:
            texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(texts)


def test_lpt_uses_project_apollo_interface(tmp_path, monkeypatch):
    prompts = []

    class Model:
        def invoke(self, prompt):
            prompts.append(prompt)
            return SimpleNamespace(content=(
                "| | Status | Next Steps |\n|---|---|---|\n"
                "| Global | Study approved | Begin recruitment |\n"
                "| China | Sites selected | Train sites |\n\n"
                "| Risk / Issue | Mitigation / Action |\n|---|---|\n"
                "| Slow recruitment | Monitor weekly |"
            ))

    path = tmp_path / "input.pptx"
    path.write_bytes(_data_pptx())
    monkeypatch.setattr(lpt, "create_apollo_llm", lambda: Model())
    bottom, source = lpt.extract_bottom_tables(path, 0)
    assert "Study approved" in bottom.global_status
    assert bottom.risks == [("Slow recruitment", "Monitor weekly")]
    assert source == "lpt_llm"
    assert "Sites selected" in prompts[0]
    with pytest.raises(ValueError, match="不存在"):
        lpt.extract_bottom_tables(path, 2)


@pytest.mark.live_llm
@pytest.mark.skipif(
    os.getenv("RUN_PROJECT_SLIDE_LIVE") != "1",
    reason="Requires explicit real Apollo acceptance run",
)
def test_real_agent_generates_modifies_and_downloads_ppt(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + (tmp_path / "sessions.db").as_posix())
    monkeypatch.setenv("AGENT_CHECKPOINT_ENABLED", "false")
    monkeypatch.setenv("AGENT_MEMORY_ENABLED", "false")
    monkeypatch.setenv("AGENT_TRUST_PROXY_IDENTITY", "false")
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("RESULT_DIR", str(tmp_path / "results"))
    from src.main import app
    from src.web.routes import agent as routes

    service = routes.AgentWebService()
    service.session_repository.initialize_schema()
    monkeypatch.setattr(routes, "agent_web_service", service)
    client = TestClient(app)
    try:
        session_id = client.post("/api/agent/sessions").json()["session_id"]
        upload = client.post(
            f"/api/agent/sessions/{session_id}/project-slides/files",
            files={"file": ("acceptance.pptx", _data_pptx())},
        )
        assert upload.status_code == 201, upload.text
        file_id = upload.json()["file_id"]
        generated = client.post(
            f"/api/agent/sessions/{session_id}/messages",
            json={"message": f"请使用上传文件生成 Project Slide 第1页并提供下载链接。[file_id: {file_id}]"},
        )
        assert generated.status_code == 200, generated.text
        result = generated.json()
        assert result["downloads"], result["answer"]
        first_url = result["downloads"][0]["url"]
        assert first_url in result["answer"]
        state = service.graph.get_state(service._config(service.sessions[session_id])).values
        assert any(item["tool_name"] == "generate_project_slide" for item in state["tool_artifacts"])
        first = client.get(first_url)
        assert first.status_code == 200
        assert "Acceptance Test Project" in _slide_text(first.content)
        from src.services.slide_generator_services import load_job_meta

        first_meta = load_job_meta(session_id, first_url.rsplit("/", 1)[1])

        modified = client.post(
            f"/api/agent/sessions/{session_id}/messages",
            json={"message": "请修改刚生成的 Project Slide，仅把标题改为 Acceptance Updated Project，其余内容保持不变，提供新下载链接。"},
        )
        assert modified.status_code == 200, modified.text
        result = modified.json()
        assert result["downloads"], result["answer"]
        second_url = result["downloads"][0]["url"]
        assert second_url != first_url
        assert second_url in result["answer"]
        state = service.graph.get_state(service._config(service.sessions[session_id])).values
        assert any(item["tool_name"] == "modify_project_slide" for item in state["tool_artifacts"])
        second = client.get(second_url)
        assert second.status_code == 200
        assert "Acceptance Updated Project" in _slide_text(second.content)
        assert first.content != second.content
        second_meta = load_job_meta(session_id, second_url.rsplit("/", 1)[1])
        for field in ("upper", "lower", "bottom"):
            assert second_meta[field] == first_meta[field]
        other_id = client.post("/api/agent/sessions").json()["session_id"]
        assert client.get(second_url.replace(session_id, other_id)).status_code == 404
        print(f"REAL_AGENT_PROJECT_SLIDE_PASSED; artifacts={tmp_path}")
    finally:
        service.session_repository.engine.dispose()