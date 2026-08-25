from pathlib import Path
from dataclasses import replace
from types import SimpleNamespace
import sys

import pytest
from openpyxl import Workbook

from src.mineru.services.mineru_client import BatchParseResult, ParseResult
from src.schemas.document_dmf import ExtractedDMFQueryBatch
from src.services.document_dmf_service import DocumentDMFError, DocumentDMFService
from src.settings import load_settings


def _settings(tmp_path: Path):
	return load_settings(
		{
			"MINERU_TOKEN": "test-token",
			"RESULT_DIR": str(tmp_path / "results"),
			"UPLOAD_DIR": str(tmp_path / "uploads"),
			"TEMP_DIR": str(tmp_path / "temp"),
		},
	)


def test_default_upload_extensions_include_excel(tmp_path: Path) -> None:
	settings = _settings(tmp_path)

	assert {".xlsx", ".xls"} <= settings.allowed_extensions
	assert {".md", ".txt"} <= settings.allowed_extensions


def test_markdown_document_is_stored_and_extracted(tmp_path: Path) -> None:
	source = tmp_path / "sample.md"
	source.write_text("原料 Ibuprofen", encoding="utf-8")
	seen: list[str] = []
	service = DocumentDMFService(
		_settings(tmp_path),
		extractor=lambda markdown: (
			seen.append(markdown)
			or {"ingredients": [" Ibuprofen ", "ibuprofen", ""]}
		),
	)

	artifact = service.parse_document(source)
	query = service.extract_query(artifact)

	assert Path(artifact.markdown_path).is_file()
	assert Path(artifact.markdown_path).is_relative_to(tmp_path / "results")
	assert seen == ["原料 Ibuprofen"]
	assert query.queries[0].ingredients == ["Ibuprofen"]


def test_service_rejects_markdown_outside_result_directory(tmp_path: Path) -> None:
	outside = tmp_path / "outside.md"
	outside.write_text("Ibuprofen", encoding="utf-8")
	service = DocumentDMFService(_settings(tmp_path), extractor=lambda text: {})

	with pytest.raises(DocumentDMFError, match="结果目录之外"):
		service.extract_query(
			{
				"document_id": "outside",
				"file_name": outside.name,
				"source_path": str(outside),
				"markdown_path": str(outside),
				"status": "parsed",
			}
		)


def test_confirmed_query_is_the_only_dmf_execution_boundary(tmp_path: Path) -> None:
	calls: list[dict] = []
	service = DocumentDMFService(
		_settings(tmp_path),
		searcher=lambda **kwargs: calls.append(kwargs) or {"success": True},
	)
	query = ExtractedDMFQueryBatch.model_validate({"ingredients": ["Ibuprofen"]})

	assert calls == []
	assert service.execute_confirmed_query(query) == {"success": True}
	assert calls == [
		{"dmf_no": "", "applicant_name": "", "ingredients": ["Ibuprofen"]}
	]


def test_batch_flow_conditions_preserve_rows_and_execute_each_query(
	tmp_path: Path,
) -> None:
	source = tmp_path / "batch.md"
	source.write_text("three rows", encoding="utf-8")
	calls: list[dict] = []
	service = DocumentDMFService(
		_settings(tmp_path),
		extractor=lambda text: {
			"queries": [
				{"dmf_no": "234", "ingredients": ["Ibuprofen"]},
				{"dmf_no": "211", "ingredients": ["NOT"]},
				{"ingredients": ["Ibanez"]},
			]
		},
		searcher=lambda **kwargs: calls.append(kwargs) or {
			"success": True,
			"query_count": 1,
			"success_count": 1,
			"failed_count": 0,
			"total_records": 1,
			"results": [{"query": kwargs, "records": [{}]}],
		},
	)

	batch = service.extract_query(service.parse_document(source))
	result = service.execute_confirmed_query(batch)

	assert calls == [
		{"dmf_no": "234", "applicant_name": "", "ingredients": ["Ibuprofen"]},
		{"dmf_no": "211", "applicant_name": "", "ingredients": ["NOT"]},
		{"dmf_no": "", "applicant_name": "", "ingredients": ["Ibanez"]},
	]
	assert result["query_count"] == 3
	assert result["success_count"] == 3
	assert result["total_records"] == 3


def test_batch_merge_counts_captcha_or_network_failure_before_query(
	tmp_path: Path,
) -> None:
	service = DocumentDMFService(
		_settings(tmp_path),
		searcher=lambda **kwargs: {
			"success": False,
			"message": "验证码获取异常：连接超时",
			"query": {
				"dmf_no": kwargs["dmf_no"],
				"applicant_name": kwargs["applicant_name"],
				"ingredients": kwargs["ingredients"],
			},
			"query_count": 0,
			"success_count": 0,
			"failed_count": 0,
			"total_records": 0,
			"results": [],
		},
	)
	batch = ExtractedDMFQueryBatch.model_validate(
		{
			"queries": [
				{"dmf_no": "234", "ingredients": ["Ibuprofen"]},
				{"ingredients": ["Ibuprofen"]},
			]
		}
	)

	result = service.execute_confirmed_query(batch)

	assert result["success"] is False
	assert result["query_count"] == 2
	assert result["success_count"] == 0
	assert result["failed_count"] == 2
	assert len(result["results"]) == 2
	assert result["results"][0]["query"] == {
		"dmf_no": "234",
		"applicant_name": "",
		"ingredient": "Ibuprofen",
	}
	assert "连接超时" in result["results"][1]["message"]


def test_empty_flow_conditions_are_rejected(tmp_path: Path) -> None:
	source = tmp_path / "sample.md"
	source.write_text("没有查询条件", encoding="utf-8")
	service = DocumentDMFService(_settings(tmp_path), extractor=lambda text: {})
	artifact = service.parse_document(source)

	with pytest.raises(DocumentDMFError, match="未提取到可用"):
		service.extract_query(artifact)


def test_pdf_is_parsed_by_mineru_and_returns_trusted_artifact(tmp_path: Path) -> None:
	settings = _settings(tmp_path)
	source = tmp_path / "sample.pdf"
	source.write_bytes(b"%PDF-test")
	markdown = settings.result_dir / "mineru-id" / "document.md"
	markdown.parent.mkdir(parents=True)
	markdown.write_text("Ibuprofen", encoding="utf-8")
	seen: list[Path] = []

	class FakeMinerU:
		def __init__(self, configured_settings) -> None:
			assert configured_settings is settings

		def __enter__(self):
			return self

		def __exit__(self, exc_type, exc, traceback) -> None:
			return None

		def parse_files(self, paths):
			seen.extend(paths)
			return BatchParseResult(
				batch_id="batch-1",
				files=(
					ParseResult(
						data_id="mineru-id",
						original_name=source.name,
						state="done",
						markdown_path=markdown,
					),
				),
			)

	service = DocumentDMFService(settings, mineru_factory=FakeMinerU)
	artifact = service.parse_document(source)

	assert seen == [source.resolve()]
	assert artifact.document_id == "mineru-id"
	assert artifact.markdown_path == str(markdown.resolve())


def test_mineru_failure_stops_document_pipeline(tmp_path: Path) -> None:
	settings = _settings(tmp_path)
	source = tmp_path / "sample.pdf"
	source.write_bytes(b"%PDF-test")

	class FailingMinerU:
		def __init__(self, configured_settings) -> None:
			pass

		def __enter__(self):
			return self

		def __exit__(self, exc_type, exc, traceback) -> None:
			return None

		def parse_files(self, paths):
			return BatchParseResult(
				batch_id="batch-1",
				files=(
					ParseResult(
						data_id="mineru-id",
						original_name=source.name,
						state="failed",
						error="MinerU 测试失败",
					),
				),
			)

	service = DocumentDMFService(settings, mineru_factory=FailingMinerU)

	with pytest.raises(DocumentDMFError, match="MinerU 测试失败"):
		service.parse_document(source)


def test_xlsx_all_nonempty_sheets_are_converted_without_mineru(
	tmp_path: Path,
) -> None:
	settings = _settings(tmp_path)
	source = tmp_path / "dmf.xlsx"
	workbook = Workbook()
	first = workbook.active
	first.title = "DMF Records"
	first.append(["Ingredient", "Applicant", "Notes"])
	first.append(["Ibuprofen", "Example|Pharma", "line 1\nline 2"])
	second = workbook.create_sheet("Numbers")
	second.append(["DMF No", "Ingredient"])
	second.append([12345, "Naproxen"])
	workbook.create_sheet("Empty")
	workbook.save(source)

	class UnexpectedMinerU:
		def __init__(self, configured_settings) -> None:
			raise AssertionError("Excel must not be sent to MinerU")

	service = DocumentDMFService(settings, mineru_factory=UnexpectedMinerU)
	artifact = service.parse_document(source)
	markdown = Path(artifact.markdown_path).read_text(encoding="utf-8")

	assert "## DMF Records" in markdown
	assert "## Numbers" in markdown
	assert "## Empty" not in markdown
	assert "Example\\|Pharma" in markdown
	assert "line 1<br>line 2" in markdown
	assert "| 12345 | Naproxen |" in markdown


def test_xlsx_normalizes_empty_and_duplicate_headers(tmp_path: Path) -> None:
	source = tmp_path / "headers.xlsx"
	workbook = Workbook()
	sheet = workbook.active
	sheet.append(["Ingredient", "Ingredient", None])
	sheet.append(["Ibuprofen", "Naproxen", "Aspirin"])
	workbook.save(source)

	artifact = DocumentDMFService(_settings(tmp_path)).parse_document(source)
	markdown = Path(artifact.markdown_path).read_text(encoding="utf-8")

	assert "| Ingredient | Ingredient 2 | Column 3 |" in markdown


def test_empty_xlsx_is_rejected(tmp_path: Path) -> None:
	source = tmp_path / "empty.xlsx"
	Workbook().save(source)

	with pytest.raises(DocumentDMFError, match="没有可读取"):
		DocumentDMFService(_settings(tmp_path)).parse_document(source)


def test_xlsx_markdown_size_limit_is_enforced_without_artifact(
	tmp_path: Path,
) -> None:
	settings = replace(_settings(tmp_path), max_markdown_size_bytes=10)
	source = tmp_path / "large.xlsx"
	workbook = Workbook()
	workbook.active.append(["Ingredient"])
	workbook.active.append(["Ibuprofen"])
	workbook.save(source)
	documents_dir = settings.result_dir / "documents"

	with pytest.raises(DocumentDMFError, match="Markdown 超过大小限制"):
		DocumentDMFService(settings).parse_document(source)

	assert not documents_dir.exists() or not any(documents_dir.iterdir())


def test_xls_uses_xlrd_and_converts_all_nonempty_sheets(
	tmp_path: Path,
	monkeypatch,
) -> None:
	source = tmp_path / "legacy.xls"
	source.write_bytes(b"legacy-xls")

	class FakeSheet:
		def __init__(self, name, rows):
			self.name = name
			self._rows = rows
			self.nrows = len(rows)
			self.ncols = max((len(row) for row in rows), default=0)

		def cell(self, row, column):
			value = self._rows[row][column] if column < len(self._rows[row]) else ""
			return SimpleNamespace(value=value, ctype=1)

	class FakeWorkbook:
		datemode = 0

		def sheets(self):
			return [
				FakeSheet("Legacy", [["Ingredient"], ["Ibuprofen"]]),
				FakeSheet("Empty", []),
			]

		def release_resources(self):
			return None

	seen: list[Path] = []
	fake_xlrd = SimpleNamespace(
		XL_CELL_DATE=3,
		open_workbook=lambda path, on_demand: seen.append(path) or FakeWorkbook(),
		xldate_as_datetime=lambda value, datemode: value,
	)
	monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

	artifact = DocumentDMFService(_settings(tmp_path)).parse_document(source)
	markdown = Path(artifact.markdown_path).read_text(encoding="utf-8")

	assert seen == [source.resolve()]
	assert "## Legacy" in markdown
	assert "Ibuprofen" in markdown
	assert "## Empty" not in markdown